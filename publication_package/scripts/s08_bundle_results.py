"""
s08_bundle_results.py

Consolidates every results file of the project into results_bundle/ (repo root):

    RESULTS_ALL.xlsx            every results table (CSV/JSON) as one sheet each, plus an INDEX sheet
    RESULTS_ALL.md              every results write-up (summary, notes, reviews, reports) in one document
    FIGURES_ALL.pdf             every figure / diagram, one per page (needs Microsoft Edge or Chrome)
    full_results_archive.zip   all original result files + raw evidence, with their folder structure

The original files are left in place: RESULTS_SUMMARY.md, the run script and the
SHA-256 manifest refer to them. Re-run after run_all_experiments.py:

    python publication_package/scripts/s08_bundle_results.py
"""

from __future__ import annotations

import json
import shutil
import subprocess
import zipfile
from pathlib import Path

import pandas as pd
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[2]
PKG = ROOT / "publication_package"
OUT = ROOT / "results_bundle"

# (file, short description); order = order in the workbook
TABLES = [
    ("publication_package/statistical_significance.csv", "Task 1: RMSE/MAPE/R2 with 95% bootstrap CIs, per model and series"),
    ("publication_package/statistical_significance_block_sensitivity.csv", "Task 1: CI sensitivity to bootstrap block length"),
    ("publication_package/statistical_seed_variance_xgboost.csv", "Task 1: XGBoost retrained with seeds 0-4"),
    ("publication_package/test_set_spans.csv", "Task 1: test-set date spans per model"),
    ("publication_package/metric_reproduction_check.csv", "Provenance: saved models vs recorded metrics"),
    ("publication_package/gate_threshold_sweep.csv", "Task 2: gate precision/recall/F1 for thresholds 0.50-0.99"),
    ("publication_package/gate_baseline_comparison.csv", "Task 2: confidence gate vs naive, random and heuristic gates"),
    ("publication_package/gate_original_scan_realized_error.csv", "Task 2: original 9-decision scan with realised error"),
    ("publication_package/prophet_calibration_before_after.csv", "Task 3: Prophet confidence before/after conformal recalibration"),
    ("publication_package/prophet_reliability_by_confidence_bucket.csv", "Task 3: Prophet error by confidence bucket"),
    ("publication_package/lstm_uncertainty_comparison.csv", "Task 4: LSTM MC-Dropout vs deep ensemble"),
    ("publication_package/lstm_seed_variance.csv", "Task 4: LSTM accuracy across seeds"),
    ("publication_package/gas_measurements.csv", "Task 5/8: measured gas per call, contract v1 vs v2"),
    ("publication_package/gas_cost_table.csv", "Task 8: USD cost per decision at several gas prices"),
    ("publication_package/gas_cost_scale.csv", "Task 8: yearly cost at scale"),
    ("publication_package/gas_cost_l2_scenarios.csv", "Task 8: L2 / sidechain cost scenarios"),
    ("publication_package/reproducibility_manifest.csv", "Task 9: SHA-256 of every output file"),
    ("publication_package/intermediate/test_predictions_xgboost.csv", "Raw: XGBoost test-set predictions + confidence"),
    ("publication_package/intermediate/test_predictions_prophet.csv", "Raw: Prophet test-set predictions + native confidence"),
    ("publication_package/intermediate/test_predictions_prophet_conformal.csv", "Raw: Prophet predictions with conformal intervals"),
    ("publication_package/intermediate/test_predictions_lstm_mc.csv", "Raw: LSTM MC-Dropout predictions"),
    ("publication_package/intermediate/test_predictions_lstm_ensemble.csv", "Raw: LSTM deep-ensemble predictions"),
    ("publication_package/intermediate/lstm_ensemble_members.csv", "Raw: per-member LSTM ensemble predictions"),
    ("docs/model_performance.csv", "Original: model performance table (Review-2)"),
    ("docs/gate_accuracy_scan.csv", "Original: 9-decision gate scan (77.8% / 22.2%)"),
    ("docs/prophet_per_series.csv", "Original: Prophet per-series metrics"),
    ("docs/lstm_per_series.csv", "Original: LSTM per-series metrics"),
    ("docs/report_data.json", "Original: report data (flattened key/value)"),
    ("models/saved_models/benchmark_summary.json", "Original: benchmark summary (flattened)"),
    ("models/saved_models/lstm/metrics.json", "Original: LSTM saved-model metrics (flattened)"),
    ("models/saved_models/prophet/metrics.json", "Original: Prophet saved-model metrics (flattened)"),
]

DOCS = [
    ("publication_package/RESULTS_SUMMARY.md", "Results summary (all numbers, before/after, limitations)"),
    ("publication_package/statistical_significance_notes.md", "Task 1 notes: statistical method"),
    ("publication_package/gate_analysis_notes.md", "Task 2 notes: decision-gate analysis"),
    ("publication_package/prophet_calibration_notes.md", "Task 3 notes: Prophet calibration"),
    ("publication_package/lstm_uncertainty_notes.md", "Task 4 notes: LSTM uncertainty"),
    ("publication_package/security_review.md", "Task 5: smart-contract security review"),
    ("publication_package/test_report.md", "Task 6: test report"),
    ("publication_package/gas_cost_analysis.md", "Task 8: gas cost analysis"),
    ("publication_package/reproducibility_check.md", "Task 9: reproducibility check"),
    ("docs/novelty_and_contribution.md", "Original: novelty and contribution"),
    ("docs/review2_report.md", "Original: Review-2 report"),
]

FIGURES = [
    ("publication_package/gate_roc_curve.svg", "Task 2: decision-gate ROC and precision-recall curves"),
    ("publication_package/class_diagram.svg", "Task 7: class diagram (as implemented)"),
    ("publication_package/sequence_diagram.svg", "Task 7: sequence diagram, one end-to-end decision"),
    ("docs/architecture_diagram.svg", "Original: architecture diagram"),
]

# everything that goes into the archive (dirs are recursive)
ARCHIVE = [
    "publication_package", "docs", "tests", "blockchain/test", "blockchain/coverage",
    "blockchain/legacy", "blockchain/contracts", "models/saved_models/benchmark_summary.json",
    "models/saved_models/lstm/metrics.json", "models/saved_models/prophet/metrics.json",
    "requirements.txt", "requirements-lock.txt",
]
ARCHIVE_SKIP_DIRS = {"__pycache__", "lstm_ensemble"}  # caches; trained models are regenerated by s04


def _flatten(obj, prefix=""):
    if isinstance(obj, dict):
        for k, v in obj.items():
            yield from _flatten(v, f"{prefix}.{k}" if prefix else str(k))
    elif isinstance(obj, list) and any(isinstance(v, (dict, list)) for v in obj):
        for i, v in enumerate(obj):
            yield from _flatten(v, f"{prefix}[{i}]")
    elif isinstance(obj, list):
        yield prefix, ", ".join(str(v) for v in obj)
    else:
        yield prefix, obj


def _read_table(rel: str) -> pd.DataFrame:
    path = ROOT / rel
    if path.suffix == ".json":
        return pd.DataFrame(list(_flatten(json.loads(path.read_text(encoding="utf-8")))), columns=["key", "value"])
    return pd.read_csv(path)


def _sheet_name(rel: str, used: set[str]) -> str:
    stem = Path(rel).stem.replace("test_predictions_", "pred_").replace("statistical_significance", "stat_sig")
    if rel.startswith("docs/"):
        stem = "orig_" + stem
    elif rel.startswith("models/"):
        stem = "orig_" + Path(rel).parent.name + "_" + stem
    name, i = stem[:31], 2
    while name in used:
        name = f"{stem[:28]}_{i}"
        i += 1
    used.add(name)
    return name


def build_workbook():
    used: set[str] = {"INDEX"}
    index = []
    path = OUT / "RESULTS_ALL.xlsx"
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        pd.DataFrame().to_excel(xw, sheet_name="INDEX")  # placeholder so INDEX is the first sheet
        for rel, desc in TABLES:
            df = _read_table(rel)
            name = _sheet_name(rel, used)
            df.to_excel(xw, sheet_name=name, index=False)
            ws = xw.sheets[name]
            ws.freeze_panes = "A2"
            for j, col in enumerate(df.columns, 1):
                width = max([len(str(col))] + [len(str(v)) for v in df[col].head(200)])
                ws.column_dimensions[get_column_letter(j)].width = min(max(width + 2, 8), 60)
            index.append({"sheet": name, "description": desc, "source_file": rel,
                          "rows": len(df), "columns": len(df.columns)})
        idx = pd.DataFrame(index)
        idx.to_excel(xw, sheet_name="INDEX", index=False)
        ws = xw.sheets["INDEX"]
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for j, w in enumerate([34, 70, 72, 8, 9], 1):
            ws.column_dimensions[get_column_letter(j)].width = w
        for r, row in enumerate(index, 2):  # make sheet names clickable
            c = ws.cell(row=r, column=1)
            c.hyperlink = f"#'{row['sheet']}'!A1"
            c.font = Font(color="0563C1", underline="single")
    return path, len(index)


def build_markdown():
    parts = ["# All results write-ups (consolidated)\n",
             "Generated by `publication_package/scripts/s08_bundle_results.py`. Each section is the "
             "unchanged text of its source file. The numbers trace to the sheets of `RESULTS_ALL.xlsx`, "
             "which carry the same file names.\n",
             "## Contents\n"]
    for i, (rel, desc) in enumerate(DOCS, 1):
        parts.append(f"{i}. {desc} (`{rel}`)")
    for i, (rel, desc) in enumerate(DOCS, 1):
        body = (ROOT / rel).read_text(encoding="utf-8").strip()
        # demote headings one level so each source file sits under its own H1
        body = "\n".join("#" + line if line.startswith("#") else line for line in body.splitlines())
        parts.append(f"\n\n---\n\n# {i}. {desc}\n\n*Source: `{rel}`*\n\n{body}")
    path = OUT / "RESULTS_ALL.md"
    path.write_text("\n".join(parts) + "\n", encoding="utf-8")
    return path


def _browser():
    for p in [r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe",
              r"C:\Program Files\Microsoft\Edge\Application\msedge.exe",
              r"C:\Program Files\Google\Chrome\Application\chrome.exe"]:
        if Path(p).exists():
            return p
    return shutil.which("msedge") or shutil.which("google-chrome") or shutil.which("chromium")


def build_figures_pdf():
    exe = _browser()
    if exe is None:
        print("(skip) FIGURES_ALL.pdf: no Edge/Chrome found")
        return None
    pages = []
    for i, (rel, desc) in enumerate(FIGURES, 1):
        pages.append(f'<section><h1>Figure {i}. {desc}</h1><p>Source: {rel}</p>'
                     f'<img src="{(ROOT / rel).as_uri()}"></section>')
    html = ("<!doctype html><meta charset='utf-8'><style>"
            "@page{size:A3 landscape;margin:10mm}"
            "body{font-family:Segoe UI,Arial,sans-serif;margin:0;color:#222}"
            "section{page-break-after:always;height:270mm;display:flex;flex-direction:column}"
            "section:last-child{page-break-after:auto}"
            "h1{font-size:16pt;margin:0 0 2mm}p{font-size:9pt;color:#666;margin:0 0 4mm}"
            "img{flex:1;min-height:0;object-fit:contain;width:100%}"
            "</style>" + "".join(pages))
    tmp = OUT / "_figures.html"
    tmp.write_text(html, encoding="utf-8")
    pdf = OUT / "FIGURES_ALL.pdf"
    profile = OUT / "_browser_profile"
    subprocess.run([exe, "--headless=new", "--disable-gpu", "--no-pdf-header-footer",
                    "--print-to-pdf-no-header", f"--user-data-dir={profile}",
                    f"--print-to-pdf={pdf}", tmp.as_uri()],
                   stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=120)
    tmp.unlink()
    shutil.rmtree(profile, ignore_errors=True)
    return pdf if pdf.exists() else None


def build_archive():
    path = OUT / "full_results_archive.zip"
    n = 0
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED, compresslevel=9) as zf:
        for rel in ARCHIVE:
            src = ROOT / rel
            files = [src] if src.is_file() else sorted(p for p in src.rglob("*") if p.is_file())
            for f in files:
                if ARCHIVE_SKIP_DIRS & set(f.relative_to(ROOT).parts):
                    continue
                zf.write(f, f.relative_to(ROOT).as_posix())
                n += 1
    return path, n


def main():
    OUT.mkdir(exist_ok=True)
    xlsx, n_sheets = build_workbook()
    md = build_markdown()
    pdf = build_figures_pdf()
    zp, n_zip = build_archive()
    print(f"{xlsx.name}: {n_sheets} tables + INDEX")
    print(f"{md.name}: {len(DOCS)} documents")
    print(f"{pdf.name if pdf else 'FIGURES_ALL.pdf (skipped)'}: {len(FIGURES)} figures")
    print(f"{zp.name}: {n_zip} files")
    for f in sorted(OUT.iterdir()):
        print(f"  {f.name:28s} {f.stat().st_size / 1024:9.1f} KB")


if __name__ == "__main__":
    main()
